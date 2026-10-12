"""Shared DOCX builder for the three IT Architecture books.

Manuscript conventions (one paragraph per line):
  # N Title            chapter heading (unnumbered '# Title' for front and back matter)
  ## Heading / ### Sub  section headings
  - item               bullet
  1. item              numbered step (rendered with a hanging indent)
  N.M text             exercise (spacing tuned for exercise lists)
  **N.M** text         answer entry
  > text               callout box; start with a bold label such as **Common mistake.**
  | a | b |            table; the first row is the header
  ~~~ ... ~~~          code listing
  FIGURE name | Caption    figure from figures/name.png
  {FIG:name}           replaced by 'Figure N'
  {REF:key}            numbered IEEE citation from references.py
"""
from pathlib import Path
import json, re, argparse
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BLUE = '276EF1'
MAX_FIG_HEIGHT = 12.5  # cm
GRAY = 'F3F3F3'
BODY_FONT = 'Nimbus Sans'
MONO_FONT = 'DejaVu Sans Mono'
ACCESSED = 'Oct. 11, 2026'
COLOPHON = ('IT Architecture Series. Revised edition, October 2026.\n'
            'All organisations, people, workloads, budgets and datasets in this book are fictional teaching examples. '
            'Any resemblance to real organisations is coincidental.\n'
            'Product and company names are trademarks of their respective owners and are used only to identify '
            'their products. This book is a study text, not professional, legal or financial advice.\n'
            'Set in Nimbus Sans and DejaVu Sans Mono.')


def set_fonts(rpr_owner, name):
    """Set every font slot explicitly so theme fonts (Calibri/Carlito) never win."""
    rPr = rpr_owner.get_or_add_rPr()
    fonts = rPr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rPr.insert(0, fonts)
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + attr), None)
    for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + attr), name)


def shade(pPr_owner, fill, left_border=None):
    pPr = pPr_owner.get_or_add_pPr()
    if left_border:
        b = OxmlElement('w:pBdr'); l = OxmlElement('w:left')
        for k, v in [('val', 'single'), ('sz', '18'), ('space', '6'), ('color', left_border)]:
            l.set(qn('w:' + k), v)
        b.append(l); pPr.append(b)
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    pPr.append(sh)


def index_terms(root):
    """Terms for the generated index: glossary entries plus multi-word terms set in bold where chapters
    define them. Single words come only from the glossary, which keeps generic words out of the index."""
    root = Path(root); terms = {}; glossary = set()
    skip_first = {'a', 'an', 'the', 'analyse', 'present', 'identify', 'brainstorm', 'generate', 'test', 'opens',
                  'fans', 'one', 'several', 'gate', 'load,', 'cash', 'design', 'option', 'assigned'}

    def norm(term):
        # Sentence-case terms become lower case; names and acronyms keep their capitals.
        base = term.split(' (')[0]
        if term[0].isupper() and base[1:2].islower() and base[1:] == base[1:].lower() and "'s" not in base:
            return term[0].lower() + term[1:]
        return term

    def key(t):
        k = t.lower()
        if k.endswith('ies') and len(k) > 5: return k[:-3] + 'y'
        if k.endswith('s') and not k.endswith('ss') and len(k) > 4: return k[:-1]
        return k

    def add(term):
        k = key(term)
        if k not in terms or len(term) < len(terms[k]): terms[k] = term

    for f in sorted((root / 'manuscript').glob('*.md')):
        if 'glossary' in f.name:
            for m in re.finditer(r'^\*\*(.+?)\.\*\*', f.read_text(), re.M):
                term = norm(m.group(1).strip()); add(term); glossary.add(key(term))
    for f in sorted((root / 'manuscript').glob('*.md')):
        if not (re.fullmatch(r'\d\d\.md', f.name) and f.name != '00.md'): continue
        t = re.sub(r'^~~~.*?^~~~[^\n]*$', '', f.read_text(), flags=re.M | re.S)
        for m in re.finditer(r'\*\*([^*\n]+?)\*\*', t):
            term = m.group(1).strip()
            words = term.split()
            if (term[-1] in '.:;,!?' or re.search(r'\d', term) or not term[0].isalpha()
                    or len(words) > 5 or len(term) > 48 or words[0].lower() in skip_first):
                continue
            term = norm(term)
            if re.match(r'^[a-z]+[A-Z]', term): continue          # code identifiers such as factOrderLine
            if len(words) == 1 and key(term) not in glossary and not term.isupper(): continue
            if key(term) in terms and key(term) in glossary: continue
            add(term)
    # Optional extra entries (one per line), e.g. author names discussed in the text.
    extra = root / 'index_extra.txt'
    if extra.exists():
        for term in extra.read_text().splitlines():
            if term.strip() and not term.startswith('#'): add(term.strip())
    return sorted(terms.values(), key=str.lower)


def defer_figures(lines, root):
    """Move a figure past the next few body paragraphs when release.py found that it left a large gap
    at the foot of the previous page (figure_defer.json: name -> paragraphs). Never past a heading,
    table, listing or another figure, so figure order and section membership are unchanged."""
    f = Path(root) / 'figure_defer.json'
    defer = json.loads(f.read_text()) if f.exists() else {}
    if not defer: return lines
    out = list(lines); i = 0
    while i < len(out):
        line = out[i]
        name = line[7:].split(' | ', 1)[0].strip() if line.startswith('FIGURE ') else None
        if name and defer.get(name):
            need, j = defer[name], i + 1
            while need and j < len(out):
                nxt = out[j]
                if not nxt.strip(): j += 1; continue
                if nxt.startswith(('#', 'FIGURE ', '~~~', '|', '>')): break
                j += 1
                # A paragraph ends at the next blank line or block; bullet runs count line by line.
                need -= 1
            if j > i + 1:
                fig = out.pop(i); out.insert(j - 1, fig)
                out.insert(j - 1, ''); i = j + 1; continue
        i += 1
    return out


def smart_quotes(s):
    """Typographic quotes and apostrophes outside `code` spans."""
    parts = re.split(r'(`[^`]+`)', s)
    out, prev = [], ' '
    for part in parts:
        if part.startswith('`') and part.endswith('`') and len(part) > 1:
            out.append(part); prev = 'x'; continue
        chars = []
        for c in part:
            if c == '"':
                c = '\u201c' if prev.isspace() or prev in '([{\u2013\u2014/' else '\u201d'
            elif c == "'":
                c = '\u2018' if prev.isspace() or prev in '([{\u2013\u2014/\u201c' else '\u2019'
            chars.append(c); prev = c
        out.append(''.join(chars))
    return ''.join(out)


def build(root, cfg):
    root = Path(root)
    import sys; sys.path.insert(0, str(root))
    from references import SOURCES
    (root / 'output').mkdir(exist_ok=True)
    ap = argparse.ArgumentParser()
    ap.add_argument('--pages', type=Path)
    args = ap.parse_args()
    files = sorted((root / 'manuscript').glob('*.md'))
    raw = '\n\n'.join(p.read_text() for p in files)

    # Figure numbers follow FIGURE line order; {FIG:name} cross-references resolve to them.
    # Structural scans ignore code listings, where '# comment' lines are common.
    def strip_code(t): return re.sub(r'^~~~.*?^~~~[^\n]*$', '', t, flags=re.M | re.S)
    fig_names = [m.group(1).strip() for m in re.finditer(r'^FIGURE ([^|]+)\|', strip_code(raw), re.M)]
    if len(set(fig_names)) != len(fig_names):
        raise ValueError('Duplicate figure names')
    fig_no = {n: i + 1 for i, n in enumerate(fig_names)}
    for n in fig_names:
        if not (root / 'figures' / f'{n}.png').exists():
            raise ValueError('Missing figure PNG ' + n)

    def figref(m):
        if m.group(1) not in fig_no: raise ValueError('Unknown figure ' + m.group(1))
        return f'Figure {fig_no[m.group(1)]}'
    raw = re.sub(r'\{FIG:([^}]+)\}', figref, raw)

    keys = []
    def cite(m):
        out = []
        for key in [k.strip() for k in m.group(1).split(',')]:
            if key not in SOURCES: raise ValueError('Missing source ' + key)
            if key not in keys: keys.append(key)
            out.append(str(keys.index(key) + 1))
        return '[' + '], ['.join(out) + ']'
    text = re.sub(r'\{REF:([^}]+)\}', cite, raw)
    assert '{REF:' not in text and '{FIG:' not in text
    unused = sorted(set(SOURCES) - set(keys))

    headings = []
    prose = strip_code(text)
    for m in re.finditer(r'^# (.+)$', prose, re.M):
        label = m.group(1)
        n = re.match(r'^(\d+) (.+)$', label)
        headings.append({'raw': label, 'title': n.group(2) if n else label, 'chapter': int(n.group(1)) if n else None})
    headings.append({'raw': 'References', 'title': 'References', 'chapter': None})
    page_map = json.loads(args.pages.read_text()) if args.pages and args.pages.exists() else {}
    index_file = root / 'index.json'
    index_data = json.loads(index_file.read_text()) if args.pages and index_file.exists() else {}
    if index_data:
        headings.append({'raw': 'Index', 'title': 'Index', 'chapter': None})

    stats = {
        'chapters': sum(1 for h in headings if h['chapter']),
        'exercises': len(re.findall(r'^\d+\.\d+ ', prose, re.M)),
        'answers': len(re.findall(r'^\*\*\d+\.\d+\*\*', prose, re.M)),
        'figures': len(fig_names),
        'labs': len(re.findall(r'^## Lab \d+', prose, re.M)),
        'references': len(keys),
    }

    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21)
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.95)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(.8); sec.footer_distance = Cm(.85)
    sec.different_first_page_header_footer = True
    styles = doc.styles
    for element in styles.element.xpath('//w:pBdr'):
        element.getparent().remove(element)
    for name in ['Normal', 'Body Text', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'Header', 'Footer']:
        s = styles[name]
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.name = BODY_FONT
        set_fonts(s.element, BODY_FONT)
        col = s.element.get_or_add_rPr().find(qn('w:color'))
        if col is not None:
            for attr in ['themeColor', 'themeTint', 'themeShade']: col.attrib.pop(qn('w:' + attr), None)
    # Document defaults also carry theme fonts; replace them.
    for rf in doc.styles.element.xpath('//w:docDefaults//w:rFonts'):
        for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            rf.attrib.pop(qn('w:' + attr), None)
        for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            rf.set(qn('w:' + attr), BODY_FONT)
    st = styles['Normal']
    st.font.size = Pt(10)
    st.paragraph_format.line_spacing = 1.0
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.widow_control = True
    styles['Title'].font.size = Pt(44)
    styles['Title'].font.bold = True
    styles['Title'].paragraph_format.space_after = Pt(18)
    styles['Title'].paragraph_format.line_spacing = 1.0
    styles['Subtitle'].font.size = Pt(14)
    styles['Subtitle'].font.italic = False
    styles['Subtitle'].font.bold = False
    styles['Heading 1'].font.size = Pt(22); styles['Heading 1'].font.bold = True
    styles['Heading 1'].paragraph_format.space_after = Pt(8)
    styles['Heading 1'].paragraph_format.space_before = Pt(0)
    styles['Heading 2'].font.size = Pt(12.5); styles['Heading 2'].font.bold = True
    styles['Heading 2'].paragraph_format.space_before = Pt(10)
    styles['Heading 2'].paragraph_format.space_after = Pt(4)
    styles['Heading 3'].font.size = Pt(10.5); styles['Heading 3'].font.bold = True
    styles['Heading 3'].paragraph_format.space_before = Pt(7)
    styles['Heading 3'].paragraph_format.space_after = Pt(2)
    styles['Caption'].font.size = Pt(9)
    styles['Caption'].font.color.rgb = RGBColor(65, 65, 65)
    styles['Caption'].font.italic = False
    styles['Caption'].paragraph_format.space_before = Pt(5)
    styles['Caption'].paragraph_format.space_after = Pt(10)
    for name in ['Heading 1', 'Heading 2', 'Heading 3']:
        styles[name].paragraph_format.keep_with_next = True
    st = styles.add_style('Code', 1)
    st.font.name = MONO_FONT; st.font.size = Pt(7.6); set_fonts(st.element, MONO_FONT)
    st.paragraph_format.line_spacing = 1
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.left_indent = Cm(.25)
    st.paragraph_format.keep_together = True
    st = styles.add_style('Reference', 1); st.base_style = styles['Normal']
    st.font.size = Pt(9.2); st.paragraph_format.line_spacing = 1.05
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.left_indent = Cm(.7); st.paragraph_format.first_line_indent = Cm(-.7)

    def field(p, code):
        r = p.add_run()
        f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), code)
        rr = OxmlElement('w:r'); tt = OxmlElement('w:t'); tt.text = '1'; rr.append(tt); f.append(rr)
        r._r.addnext(f)

    def bookmark(p, name, idx):
        start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(idx)); start.set(qn('w:name'), name)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(idx))
        p._p.insert(0, start); p._p.append(end)

    def link(p, label, anchor=None, url=None, size=10.3):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        h = OxmlElement('w:hyperlink')
        if anchor: h.set(qn('w:anchor'), anchor)
        else: h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
        co = OxmlElement('w:color'); co.set(qn('w:val'), '000000'); rp.append(co)
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(round(size * 2))); rp.append(sz)
        r.append(rp); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)

    def inline(p, s, size=None):
        # **bold**, *italic* and `code` spans; everything else is plain text.
        s = smart_quotes(s)
        for j, part in enumerate(re.split(r'\*\*(.*?)\*\*', s)):
            for k, sub in enumerate(re.split(r'`([^`]+)`', part)):
                if not sub: continue
                pieces = [sub] if k % 2 == 1 else re.split(r'(?<![*\w])\*([^*\n]+?)\*(?![*\w])', sub)
                for m, piece in enumerate(pieces):
                    if not piece: continue
                    r = p.add_run(piece); r.bold = (j % 2 == 1); r.italic = (m % 2 == 1) or None
                    if k % 2 == 1:
                        r.font.name = MONO_FONT; set_fonts(r._r, MONO_FONT); r.font.size = Pt(size - 1.2 if size else 8.6)
                    elif size:
                        r.font.size = Pt(size)

    hp = sec.header.paragraphs[0]
    hp.style = styles['Caption']; hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(cfg['header'] + '\t'); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0, 0, 0); r.bold = True
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), ' STYLEREF "Heading 1" ')
    rr = OxmlElement('w:r'); rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '16'); rpr.append(sz)
    co = OxmlElement('w:color'); co.set(qn('w:val'), '414141'); rpr.append(co); rr.append(rpr)
    tt = OxmlElement('w:t'); tt.text = ' '; rr.append(tt); f.append(rr); hp._p.append(f)
    fp = sec.footer.paragraphs[0]
    styles['Footer'].paragraph_format.tab_stops.clear_all()
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run(f"BOOK {cfg['number']:02d}\t"); r.font.size = Pt(8)
    field(fp, 'PAGE')

    # Cover. Counts are computed from the manuscript so they cannot drift.
    p = doc.add_paragraph(f"{cfg['number']:02d}")
    p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(30)
    for r in p.runs: r.font.size = Pt(86); r.font.color.rgb = RGBColor.from_string(BLUE)
    doc.add_paragraph(cfg['title'], 'Title')
    p = doc.add_paragraph(cfg['subtitle'], 'Subtitle'); p.paragraph_format.space_after = Pt(36)
    for s in cfg['cover_lines'](stats):
        p = doc.add_paragraph(s); p.paragraph_format.space_after = Pt(10)
        for r in p.runs: r.font.size = Pt(11)
    p = doc.add_paragraph(cfg['edition'])
    p.paragraph_format.space_before = Pt(28)
    for r in p.runs: r.font.size = Pt(10)
    p = doc.add_paragraph(cfg.get('colophon', COLOPHON))
    p.paragraph_format.space_before = Pt(cfg.get('colophon_space', 150)); p.paragraph_format.line_spacing = 1.1
    for r in p.runs: r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(65, 65, 65)
    doc.add_page_break()

    p = doc.add_paragraph('Contents', 'Heading 1')
    bookmark(p, 'contents', 1)
    for i, h in enumerate(headings):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)
        label = (str(h['chapter']).zfill(2) + '  ' if h['chapter'] else '') + h['title']
        link(p, label, anchor=f'section_{i}', size=10.1)
        r = p.add_run('\t' + str(page_map.get(h['title'], ''))); r.font.size = Pt(10.1); r.font.color.rgb = RGBColor.from_string(BLUE)

    def new_heading(i):
        h = headings[i]
        if h['chapter']:
            p = doc.add_paragraph('CHAPTER ' + str(h['chapter']))
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs: r.font.size = Pt(9); r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
        p = doc.add_paragraph(h['title'], 'Heading 1')
        if not h['chapter']:
            # Front matter follows the contents; back matter (labs, answers, glossary,
            # references) starts on a fresh page so it is easy to find.
            p.paragraph_format.page_break_before = True
        bookmark(p, f'section_{i}', i + 10)

    def table(rows):
        n = len(rows[0])
        if any(len(r) != n for r in rows):
            raise ValueError('Ragged table starting ' + ' | '.join(rows[0]))
        # Width follows content length, with a floor so short columns stay readable.
        lens = [max(4, sum(len(r[c]) for r in rows) / len(rows)) ** .75 for c in range(n)]
        widths = [16.6 * l / sum(lens) for l in lens]
        # No column may be narrower than its longest word (about 0.18 cm per character at
        # 8.6 pt bold, plus cell padding); otherwise words break mid-word.
        floor = [max(1.3, 0.35 + 0.18 * max(len(w) for r in rows for w in (re.sub(r'[*`]', '', r[c]).split() or [''])))
                 for c in range(n)]
        for _ in range(5):
            widths = [max(f, w) for f, w in zip(floor, widths)]
            excess = sum(widths) - 16.6
            if excess <= 0.01: break
            flex = [w - f for f, w in zip(floor, widths)]
            if sum(flex) <= 0: break
            widths = [w - excess * fl / sum(flex) for w, fl in zip(widths, flex)]
        scale = 16.6 / sum(widths); widths = [w * scale for w in widths]
        t = doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        for col, w in zip(t.columns, widths): col.width = Cm(w)
        for rowidx, row in enumerate(rows):
            cells = t.rows[0].cells if rowidx == 0 else t.add_row().cells
            trPr = t.rows[rowidx]._tr.get_or_add_trPr()
            trPr.append(OxmlElement('w:cantSplit'))
            if rowidx == 0: trPr.append(OxmlElement('w:tblHeader'))
            for c, w, txt in zip(cells, widths, row):
                c.width = Cm(w); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = c._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for edge in ['top', 'left', 'bottom', 'right']:
                    e = OxmlElement('w:' + edge); e.set(qn('w:w'), '85'); e.set(qn('w:type'), 'dxa'); margins.append(e)
                pr.append(margins)
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '000000' if rowidx == 0 else (GRAY if rowidx % 2 == 0 else 'FFFFFF')); pr.append(sh)
                borders = OxmlElement('w:tcBorders')
                for edge in ['top', 'left', 'bottom', 'right']:
                    el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
                pr.append(borders)
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.02
                if len(rows) <= 15 and rowidx < len(rows) - 1: p.paragraph_format.keep_with_next = True
                inline(p, txt, size=8.6)
                for r in p.runs:
                    if rowidx == 0: r.bold = True
                    r.font.color.rgb = RGBColor.from_string('FFFFFF' if rowidx == 0 else '000000')
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1; p.add_run().font.size = Pt(2)

    lines = defer_figures(text.splitlines(), root)
    fs = Path(root) / 'figure_scale.json'
    fig_scale = json.loads(fs.read_text()) if fs.exists() else {}
    i = 0; section_idx = 0; figidx = 0
    opening = 0  # >0 while inside a chapter opening that must stay with what follows
    last_para = None
    while i < len(lines):
        line = lines[i]
        if not line.strip(): i += 1; continue
        if last_para is not None and (line.startswith('~~~') or line.startswith('|')):
            # A listing's introduction always stays with it; a table's only when it ends in a colon.
            if line.startswith('~~~') or last_para.text.rstrip().endswith(':'):
                last_para.paragraph_format.keep_with_next = True
        last_para = None
        if line.startswith('~~~'):
            i += 1; block = []
            while i < len(lines) and not lines[i].startswith('~~~'):
                block.append(lines[i]); i += 1
            for j, s in enumerate(block):
                if len(s) > 92: raise ValueError('Code line longer than 92 characters: ' + s)
                p = doc.add_paragraph(s if s else ' ', 'Code')
                p.paragraph_format.keep_with_next = j < len(block) - 1 and len(block) <= 45
                shade(p._p, GRAY)
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.add_run().font.size = Pt(2)
            i += 1; continue
        if line.startswith('# '):
            new_heading(section_idx); section_idx += 1; i += 1
            opening = 2 if headings[section_idx - 1]['chapter'] else 0
            continue
        if line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 2'); i += 1
            if opening: opening -= 1
            continue
        if line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 3'); i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [s.strip() for s in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[: -]+', c) for c in cells): rows.append(cells)
                i += 1
            table(rows); continue
        if line.startswith('FIGURE '):
            name, caption = line[7:].split(' | ', 1); figidx += 1
            p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            png = root / 'figures' / f'{name.strip()}.png'
            with open(png, 'rb') as fh: head = fh.read(24)
            pw, ph = int.from_bytes(head[16:20], 'big'), int.from_bytes(head[20:24], 'big')
            width = cfg.get('figure_widths', {}).get(name.strip(), 16.3)
            if width * ph / pw > MAX_FIG_HEIGHT: width = MAX_FIG_HEIGHT * pw / ph
            width *= fig_scale.get(name.strip(), 1)  # set by release.py to close a page gap
            shape = p.add_run().add_picture(str(png), width=Cm(width))
            docPr = shape._inline.docPr; docPr.set('descr', caption); docPr.set('title', caption)
            doc.add_paragraph(f'Figure {figidx}  {caption}', 'Caption')
            i += 1; continue
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(.3); p.paragraph_format.right_indent = Cm(.2)
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_together = True
            shade(p._p, GRAY, left_border=BLUE)
            inline(p, line[2:], size=9.5)
            i += 1; continue
        p = doc.add_paragraph()
        if line.startswith('- '):
            p.paragraph_format.left_indent = Cm(.5)
            p.paragraph_format.first_line_indent = Cm(-.35)
            p.paragraph_format.space_after = Pt(1.5)
            line = '•  ' + line[2:]
        elif re.match(r'^\d+\. ', line):
            p.paragraph_format.left_indent = Cm(.6)
            p.paragraph_format.first_line_indent = Cm(-.45)
            p.paragraph_format.space_after = Pt(1.5)
        elif re.match(r'^(\*\*)?\d+\.\d+(\*\*)? ', line):
            p.paragraph_format.space_after = Pt(4)
            # Exercise numbers are bold like answer numbers.
            line = re.sub(r'^(\d+\.\d+) ', r'**\1** ', line)
        if opening: p.paragraph_format.keep_with_next = True
        inline(p, line)
        last_para = p
        i += 1

    new_heading(section_idx)
    doc.add_paragraph('References are numbered in order of first citation. Online sources were accessed in October 2026. Product documentation changes; check it again before implementation.')
    for n, key in enumerate(keys, 1):
        author, title, url, detail = SOURCES[key]
        p = doc.add_paragraph(style='Reference')
        tail = (detail.rstrip('.') + '. ' if detail else '')
        # IEEE: books and standards have italic titles; articles and web pages are quoted.
        book = bool(detail) and 'working paper' not in detail and not detail.startswith('in ') and bool(re.search(
            r"ISBN|\b\d+(st|nd|rd|th) ed\.|Press|Springer|Technics|^ISO|^ISO/IEC", detail))
        ttl = f'*{title}*, ' if book else f'"{title}," '
        if url and 'doi.org/' in url:
            # The DOI link replaces any DOI already written in the detail text.
            d = re.sub(r',?\s*doi:\s*\S+$', '', detail.rstrip('.')) if detail else ''
            tail = d + '. ' if d else ''
            inline(p, f'[{n}] {author}, {ttl}{tail}doi: ')
            link(p, url.split('doi.org/', 1)[1], url=url, size=9.2)
        elif url:
            inline(p, f'[{n}] {author}, {ttl}{tail}Accessed: {ACCESSED}. [Online]. Available: ')
            # Allow line breaks after '/' and '-' in the path, never inside the scheme's '//'.
            scheme, rest = url.split('://', 1) if '://' in url else ('', url)
            shown = (scheme + '://' if scheme else '') + rest.replace('/', '/\u200b').replace('-', '-\u200b')
            link(p, shown, url=url, size=9.2)
        else:
            inline(p, f'[{n}] {author}, ' + (f'*{title}*. ' if book else f'"{title}." ') + tail.rstrip())
        bookmark(p, f'ref_{n}', 500 + n)
    if index_data:
        from docx.enum.section import WD_SECTION
        new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
        cols = new_sec._sectPr.xpath('./w:cols')
        c = cols[0] if cols else OxmlElement('w:cols')
        c.set(qn('w:num'), '2'); c.set(qn('w:space'), '567')
        if not cols: new_sec._sectPr.append(c)
        p = doc.add_paragraph('Index', 'Heading 1'); bookmark(p, f'section_{len(headings) - 1}', len(headings) + 9)
        p = doc.add_paragraph('Page numbers point to definitions and main discussions in the chapters'
                              + (' and labs.' if stats['labs'] else '.'))
        for r in p.runs: r.font.size = Pt(8.5)
        letter = None
        for term in sorted(index_data, key=lambda t: (t[0].isalpha(), re.sub(r'[^a-z0-9 ]', '', t.lower()))):
            first = term[0].upper() if term[0].isalpha() else '0–9'
            if first != letter:
                letter = first
                p = doc.add_paragraph(letter); p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
                for r in p.runs: r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string(BLUE)
            pages = sorted(index_data[term]); parts = []; start = prev = None
            for pg in pages + [None]:
                if start is None: start = prev = pg; continue
                if pg is not None and pg == prev + 1: prev = pg; continue
                parts.append(str(start) if start == prev else f'{start}–{prev}'); start = prev = pg
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0.5)
            p.paragraph_format.left_indent = Cm(.35); p.paragraph_format.first_line_indent = Cm(-.35)
            r = p.add_run(term + '  '); r.font.size = Pt(8.5)
            r = p.add_run(', '.join(parts)); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(65, 65, 65)
    settings = doc.settings.element
    uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true'); settings.append(uf)
    doc.core_properties.title = cfg['header']
    doc.core_properties.subject = re.sub(r'\s+', ' ', cfg['subtitle'])
    doc.core_properties.author = 'IT Architecture Series'
    doc.core_properties.keywords = 'IT-architecture; software-architecture; solution-architecture; enterprise-architecture; textbook'
    out = root / 'output' / cfg['output']
    # Headings, captions and contents entries are written directly; give them typographic quotes too.
    for para in doc.paragraphs:
        if para.style.name != 'Code':
            for r in para.runs:
                if ('"' in r.text or "'" in r.text) and r.font.name != MONO_FONT: r.text = smart_quotes(r.text)
    doc.save(out)
    (root / 'heading_index.json').write_text(json.dumps(headings, indent=2))
    (root / 'reference_order.json').write_text(json.dumps([{'number': i + 1, 'key': k, 'source': SOURCES[k]} for i, k in enumerate(keys)], indent=2))
    print(json.dumps({'file': str(out), 'words': len(text.split()), **stats, 'unused_sources': unused}))
    if stats['exercises'] != stats['answers']:
        print(f"WARNING: {stats['exercises']} exercises but {stats['answers']} answers")
